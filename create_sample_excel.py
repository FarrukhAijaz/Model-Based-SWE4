"""
Generate a sample test data Excel file for the SimulinkSample model.

The model has:
  Inputs:  Throttle (double), Pedal (double)
  Outputs: Speed, Status (uint8)

Run this script to create sample_test_data.xlsx:
    python create_sample_excel.py
"""

import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "TestData"

# Headers — must match model Inport names + Expected_<Outport>
headers = ["Throttle", "Pedal", "Expected_Speed", "Expected_Status"]
ws.append(headers)

# Style header row
from openpyxl.styles import Font, PatternFill, Alignment

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center")

# Sample test data rows
test_rows = [
    # Throttle, Pedal, Expected_Speed, Expected_Status
    [0,     0,    0.0,    0],
    [25,    10,   20.5,   1],
    [50,    30,   45.0,   1],
    [75,    60,   70.5,   1],
    [100,   80,   90.0,   2],
    [100,   100,  100.0,  2],
    [50,    0,    30.0,   1],
    [0,     100,  10.0,   0],
]

for row in test_rows:
    ws.append(row)

# Auto-width columns
for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    ws.column_dimensions[col[0].column_letter].width = max_len + 4

wb.save("sample_test_data.xlsx")
print("Created: sample_test_data.xlsx")
print(f"  {len(test_rows)} test rows with columns: {headers}")
